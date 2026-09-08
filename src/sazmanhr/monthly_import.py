"""Controlled monthly XLSX import for the locked SazmanHR v4.9 frontend.

The browser always performs a dry-run first.  The preview is immutable and is
applied only after an explicit confirmation.  Blank spreadsheet cells preserve
existing values unless the field is explicitly named in ``clear_fields``.
Organizational changes are never allowed to bypass personnel movement history.
"""
from __future__ import annotations

import hashlib
import io
import json
import secrets
from datetime import datetime, timezone
from typing import Any

from openpyxl import load_workbook

from .compat_v49 import load_dataset
from .database import Repository, canonical, utc_now

MAX_IMPORT_BYTES = 24 * 1024 * 1024
PREVIEW_TTL_MINUTES = 30

PERSONNEL_HEADERS = (
    "personnel_no", "action", "full_name", "employment_group", "employment_subtype",
    "status", "position_title", "position_code", "organizational_unit", "actual_location",
    "activity_area", "company", "clear_fields", "note",
)
ASSIGNMENT_HEADERS = ("personnel_no", "assignment_type", "target_page", "target_node_id", "reason")
PERSONNEL_FIELDS = (
    "full_name", "employment_group", "employment_subtype", "status", "position_title",
    "position_code", "organizational_unit", "actual_location", "activity_area", "company",
)
ORG_FIELDS = frozenset({"position_title", "position_code", "organizational_unit", "actual_location"})
FIELD_LABELS = {
    "full_name": "نام و نام خانوادگی",
    "employment_group": "نوع همکاری",
    "employment_subtype": "زیرنوع همکاری",
    "status": "وضعیت",
    "position_title": "عنوان پست",
    "position_code": "کد پست",
    "organizational_unit": "واحد سازمانی",
    "actual_location": "محل خدمت واقعی",
    "activity_area": "حوزه فعالیت",
    "company": "شرکت",
}
ALLOWED_ACTIONS = frozenset({"update", "insert", "upsert"})
ALLOWED_ASSIGNMENT_TYPES = frozenset({"current", "approved_transfer"})


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _header_map(ws: Any, expected: tuple[str, ...]) -> dict[str, int]:
    row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    actual = [_text(value) for value in row]
    mapping = {name: actual.index(name) for name in expected if name in actual}
    missing = [name for name in expected if name not in mapping]
    if missing:
        raise ValueError(f"ستون‌های الزامی Sheet «{ws.title}» ناقص است: {', '.join(missing)}")
    return mapping


def _split_clear_fields(raw: str) -> set[str]:
    normalized = raw.replace("،", ",").replace(";", ",")
    return {part.strip() for part in normalized.split(",") if part.strip()}


def _load_position_catalog(repo: Repository) -> dict[str, dict[str, Any]]:
    try:
        payload = load_dataset(repo, "position-catalog")
    except KeyError as exc:
        raise ValueError("فهرست معتبر پست‌ها روی سرور بارگذاری نشده است.") from exc
    positions = payload.get("positions") if isinstance(payload, dict) else None
    if not isinstance(positions, list):
        raise ValueError("فهرست پست‌های سرور معتبر نیست.")
    return {str(item.get("id", "")): item for item in positions if isinstance(item, dict) and item.get("id")}


def _person_rows(repo: Repository) -> tuple[dict[str, dict[str, Any]], dict[str, dict[str, Any]]]:
    with repo.connect() as conn:
        rows = [dict(r) for r in conn.execute("SELECT * FROM personnel ORDER BY personnel_no,id").fetchall()]
    by_no = {str(r["personnel_no"]): r for r in rows}
    by_id = {str(r["id"]): r for r in rows}
    return by_no, by_id


def preview_xlsx(repo: Repository, raw: bytes, filename: str) -> tuple[dict[str, Any], dict[str, Any]]:
    if not raw:
        raise ValueError("فایل Excel خالی است.")
    if len(raw) > MAX_IMPORT_BYTES:
        raise ValueError("حجم فایل Excel بیش از ۲۴ مگابایت است.")
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("فقط فایل استاندارد .xlsx پذیرفته می‌شود.")
    file_sha = hashlib.sha256(raw).hexdigest()
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("فایل Excel قابل خواندن نیست.") from exc
    required_sheets = {"Personnel", "Assignments"}
    missing_sheets = sorted(required_sheets - set(wb.sheetnames))
    if missing_sheets:
        raise ValueError("Sheetهای الزامی فایل موجود نیست: " + "، ".join(missing_sheets))

    people_by_no, _ = _person_rows(repo)
    catalog = _load_position_catalog(repo)
    warnings: list[str] = []
    errors: list[str] = []
    seen_personnel: set[str] = set()
    personnel_plan: list[dict[str, Any]] = []
    changes: list[dict[str, Any]] = []

    ws = wb["Personnel"]
    h = _header_map(ws, PERSONNEL_HEADERS)
    for excel_row, cells in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        values = {name: _text(cells[idx]) if idx < len(cells) else "" for name, idx in h.items()}
        if not any(values.values()):
            continue
        no = values["personnel_no"]
        action = values["action"].lower()
        if not no:
            errors.append(f"Personnel ردیف {excel_row}: personnel_no خالی است.")
            continue
        if no in seen_personnel:
            errors.append(f"Personnel: شماره پرسنلی {no} تکراری است.")
            continue
        seen_personnel.add(no)
        if action not in ALLOWED_ACTIONS:
            errors.append(f"Personnel {no}: action باید update، insert یا upsert باشد.")
            continue
        old = people_by_no.get(no)
        if action == "insert" and old:
            errors.append(f"Personnel {no}: action=insert ولی فرد از قبل وجود دارد.")
            continue
        if action == "update" and not old:
            errors.append(f"Personnel {no}: action=update ولی فرد در سامانه وجود ندارد.")
            continue
        clear_fields = _split_clear_fields(values["clear_fields"])
        unknown_clear = sorted(clear_fields - set(PERSONNEL_FIELDS))
        if unknown_clear:
            errors.append(f"Personnel {no}: clear_fields نامعتبر: {', '.join(unknown_clear)}")
            continue
        is_new = old is None
        new_values: dict[str, str] = {}
        diffs: list[dict[str, str]] = []
        for field in PERSONNEL_FIELDS:
            old_value = _text(old.get(field, "")) if old else ""
            incoming = values[field]
            if field in clear_fields:
                target = ""
            elif incoming:
                target = incoming
            else:
                target = old_value
            new_values[field] = target
            if target != old_value:
                diffs.append({"field": field, "label": FIELD_LABELS[field], "old": old_value, "new": target})
        if is_new and not new_values["full_name"]:
            errors.append(f"Personnel {no}: برای نیروی جدید full_name الزامی است.")
            continue
        personnel_plan.append({
            "personnel_no": no, "action": action, "is_new": is_new,
            "person_id": str(old["id"]) if old else "",
            "values": new_values, "changed_fields": [d["field"] for d in diffs],
            "note": values["note"], "excel_row": excel_row,
        })
        changes.append({
            "personnel_no": no,
            "full_name": new_values["full_name"] or (old.get("full_name", "") if old else ""),
            "is_new": is_new,
            "fields": diffs,
            "assignment": None,
        })

    assignment_by_no: dict[str, dict[str, Any]] = {}
    ws_a = wb["Assignments"]
    ah = _header_map(ws_a, ASSIGNMENT_HEADERS)
    for excel_row, cells in enumerate(ws_a.iter_rows(min_row=2, values_only=True), start=2):
        values = {name: _text(cells[idx]) if idx < len(cells) else "" for name, idx in ah.items()}
        if not any(values.values()):
            continue
        no = values["personnel_no"]
        if not no:
            errors.append(f"Assignments ردیف {excel_row}: personnel_no خالی است.")
            continue
        if no in assignment_by_no:
            errors.append(f"Assignments: برای شماره پرسنلی {no} بیش از یک جانمایی ثبت شده است.")
            continue
        kind = values["assignment_type"].lower()
        if kind not in ALLOWED_ASSIGNMENT_TYPES:
            errors.append(f"Assignments {no}: assignment_type باید current یا approved_transfer باشد.")
            continue
        node_id = values["target_node_id"]
        target = catalog.get(node_id)
        if not target:
            errors.append(f"Assignments {no}: مقصد {node_id or 'خالی'} در فهرست معتبر سرور وجود ندارد.")
            continue
        catalog_page = int(target.get("page") or 0)
        target_page = int(values["target_page"]) if values["target_page"].isdigit() else catalog_page
        if values["target_page"] and target_page != catalog_page:
            errors.append(f"Assignments {no}: صفحه مقصد با شناسه {node_id} سازگار نیست.")
            continue
        if kind == "approved_transfer" and not bool(target.get("approved")):
            errors.append(f"Assignments {no}: مقصد {node_id} برای انتقال پست مصوب مجاز نیست.")
            continue
        if kind == "approved_transfer" and bool(target.get("is_unit")):
            errors.append(f"Assignments {no}: مقصد {node_id} یک واحد سازمانی است، نه پست مصوب قابل انتصاب.")
            continue
        if no not in people_by_no and no not in seen_personnel:
            errors.append(f"Assignments {no}: فرد نه در سامانه و نه در Sheet Personnel وجود دارد.")
            continue
        old = people_by_no.get(no)
        old_extra: dict[str, Any] = {}
        if old:
            try:
                old_extra = json.loads(old.get("extra_json") or "{}")
            except (TypeError, json.JSONDecodeError):
                old_extra = {}
        old_node = _text(old_extra.get("target_node_id") or old.get("chart_node_id") if old else "")
        old_page = int(old_extra.get("current_chart_page_v421") or old.get("chart_page_no") or 0) if old else 0
        assignment = {
            "personnel_no": no, "assignment_type": kind, "target_page": target_page,
            "target_node_id": node_id, "reason": values["reason"],
            "target_title": _text(target.get("title")), "parent_title": _text(target.get("parent_title")),
            "approved": bool(target.get("approved")), "excel_row": excel_row,
            "old_page": old_page, "old_node_id": old_node,
        }
        assignment_by_no[no] = assignment
        display = next((c for c in changes if c["personnel_no"] == no), None)
        if display is None:
            display = {
                "personnel_no": no,
                "full_name": _text(old.get("full_name", "")) if old else "",
                "is_new": old is None, "fields": [], "assignment": None,
            }
            changes.append(display)
        display["assignment"] = {
            "old_title": _text(old.get("position_title", "")) if old else "",
            "old_page": old_page, "old_node_id": old_node,
            "type": kind, "new_title": assignment["target_title"],
            "new_page": target_page, "new_node_id": node_id, "reason": values["reason"],
        }

    # Central-server safety rule: any organizational field change on an existing
    # person must carry an explicit assignment row so movement history is never
    # bypassed by the Personnel sheet.
    for item in personnel_plan:
        if not item["is_new"] and ORG_FIELDS.intersection(item["changed_fields"]) and item["personnel_no"] not in assignment_by_no:
            errors.append(
                f"Personnel {item['personnel_no']}: تغییر پست/واحد/محل خدمت نیازمند ردیف متناظر در Sheet Assignments است."
            )

    plan_people = {item["personnel_no"]: item for item in personnel_plan}
    # Assignment-only existing people still need a plan entry at apply time.
    for no, assignment in assignment_by_no.items():
        if no not in plan_people:
            old = people_by_no.get(no)
            if old:
                personnel_plan.append({
                    "personnel_no": no, "action": "update", "is_new": False,
                    "person_id": str(old["id"]),
                    "values": {field: _text(old.get(field, "")) for field in PERSONNEL_FIELDS},
                    "changed_fields": [], "note": "", "excel_row": 0,
                })

    new_count = sum(1 for item in personnel_plan if item["is_new"])
    updated_count = sum(1 for item in personnel_plan if not item["is_new"] and item["changed_fields"])
    assignment_count = len(assignment_by_no)
    if not personnel_plan and not assignment_by_no:
        warnings.append("فایل هیچ ردیف قابل اعمالی ندارد.")
    summary = {
        "personnel_rows": len(seen_personnel),
        "new_people": new_count,
        "updated_people": updated_count,
        "assignment_moves": assignment_count,
        "warnings": len(warnings),
        "errors": len(errors),
    }
    preview_id = "preview-" + secrets.token_hex(16)
    preview = {
        "preview_id": preview_id,
        "filename": filename,
        "file_sha256": file_sha,
        "sheets": list(wb.sheetnames),
        "summary": summary,
        "errors": errors,
        "warnings": warnings,
        "changes": changes,
        "can_apply": not errors and bool(personnel_plan or assignment_by_no),
        "expires_in_minutes": PREVIEW_TTL_MINUTES,
    }
    plan = {
        "preview_id": preview_id, "filename": filename, "sha256": file_sha,
        "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "people": personnel_plan, "assignments": list(assignment_by_no.values()),
        "summary": summary,
    }
    return preview, plan


def apply_plan(repo: Repository, plan: dict[str, Any], actor_id: str, *, backup_filename: str = "") -> dict[str, Any]:
    """Apply an already validated plan in one SQLite write transaction."""
    now = utc_now()
    effective_date = now[:10]
    batch_id = "monthly-" + secrets.token_hex(12)
    assignment_by_no = {str(a["personnel_no"]): a for a in plan.get("assignments", [])}
    new_people = 0
    updated_people = 0
    assignments = 0

    with repo.write() as conn:
        for item in plan.get("people", []):
            no = str(item["personnel_no"])
            old = conn.execute("SELECT * FROM personnel WHERE personnel_no=?", (no,)).fetchone()
            is_new = old is None
            values = dict(item.get("values") or {})
            assignment = assignment_by_no.get(no)

            if is_new:
                if not item.get("is_new"):
                    raise ValueError(f"Personnel {no}: وضعیت داده از زمان Dry Run تغییر کرده است.")
                person_id = "person-" + secrets.token_hex(16)
                full_name = _text(values.get("full_name"))
                if not full_name:
                    raise ValueError(f"Personnel {no}: نام کامل برای نیروی جدید الزامی است.")
                extra: dict[str, Any] = {"monthly_import_batch": batch_id}
                chart_page = None
                chart_node = ""
                if assignment:
                    chart_page = int(assignment["target_page"])
                    chart_node = str(assignment["target_node_id"])
                    extra.update({
                        "target_node_id": chart_node,
                        "current_chart_page_v421": chart_page,
                        "current_chart_page_v422": chart_page,
                    })
                    if assignment["assignment_type"] == "approved_transfer":
                        extra.update({"approved_node_id": chart_node, "home_node_id": chart_node, "home_page": chart_page})
                row_values = {
                    "organizational_unit": _text(values.get("organizational_unit")),
                    "position_code": _text(values.get("position_code")),
                    "position_title": (_text(assignment.get("target_title"))
                                       if assignment and assignment.get("assignment_type") == "approved_transfer"
                                       else _text(values.get("position_title"))),
                    "employment_group": _text(values.get("employment_group")),
                    "employment_subtype": _text(values.get("employment_subtype")),
                    "status": _text(values.get("status")) or "شاغل",
                    "activity_area": _text(values.get("activity_area")),
                    "actual_location": _text(values.get("actual_location")),
                    "company": _text(values.get("company")),
                }
                conn.execute(
                    """INSERT INTO personnel(id,personnel_no,first_name,last_name,full_name,gender,
                       organizational_unit,position_code,position_title,employment_group,employment_subtype,status,
                       activity_area,actual_location,company,chart_node_id,chart_page_no,extra_json,row_version,updated_at,updated_by)
                       VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (person_id, no, full_name, "", full_name, "", row_values["organizational_unit"],
                     row_values["position_code"], row_values["position_title"], row_values["employment_group"],
                     row_values["employment_subtype"], row_values["status"], row_values["activity_area"],
                     row_values["actual_location"], row_values["company"], chart_node, chart_page,
                     canonical(extra), 1, now, actor_id),
                )
                from_id, to_id = repo._sync_person_projection(conn, person_id, row_values, actor_id, now,
                                                               effective_date=effective_date)
                after = dict(conn.execute("SELECT * FROM personnel WHERE id=?", (person_id,)).fetchone())
                repo._record(conn, actor_id, "create", "personnel", person_id, None, after, 1)
                if assignment:
                    movement_id = "movement-" + secrets.token_hex(16)
                    movement_after = dict(after)
                    conn.execute(
                        """INSERT INTO personnel_movements
                           (id,person_id,movement_type,effective_date,order_no,order_date,reason,note,
                            from_assignment_id,to_assignment_id,before_json,after_json,created_at,created_by,row_version)
                           VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,1)""",
                        (movement_id, person_id, "appointment", effective_date, "", "",
                         _text(assignment.get("reason")) or "ورود ماهانه Excel", _text(item.get("note")),
                         from_id, to_id, canonical({}), canonical(movement_after), now, actor_id),
                    )
                    repo._record(conn, actor_id, "movement", "personnel_movement", movement_id, None, movement_after, 1)
                new_people += 1
            else:
                if item.get("is_new"):
                    raise ValueError(f"Personnel {no}: وضعیت داده از زمان Dry Run تغییر کرده است.")
                person_id = str(old["id"])
                try:
                    extra = json.loads(old["extra_json"] or "{}")
                except (TypeError, json.JSONDecodeError):
                    extra = {}
                if not isinstance(extra, dict):
                    extra = {}
                before = dict(old)
                current_values = {field: _text(old[field]) for field in PERSONNEL_FIELDS if field != "full_name"}
                for field in PERSONNEL_FIELDS:
                    if field == "full_name":
                        continue
                    current_values[field] = _text(values.get(field, old[field]))
                full_name = _text(values.get("full_name", old["full_name"]))
                chart_page = old["chart_page_no"]
                chart_node = _text(old["chart_node_id"])
                if assignment and assignment.get("assignment_type") == "approved_transfer":
                    current_values["position_title"] = _text(assignment.get("target_title")) or current_values["position_title"]
                if assignment:
                    chart_page = int(assignment["target_page"])
                    chart_node = str(assignment["target_node_id"])
                    extra.update({
                        "target_node_id": chart_node,
                        "current_chart_page_v421": chart_page,
                        "current_chart_page_v422": chart_page,
                        "monthly_import_batch": batch_id,
                    })
                    if assignment["assignment_type"] == "approved_transfer":
                        extra.update({"approved_node_id": chart_node, "home_node_id": chart_node, "home_page": chart_page})
                changed = bool(item.get("changed_fields")) or bool(assignment)
                if not changed:
                    continue
                new_version = int(old["row_version"]) + 1
                conn.execute(
                    """UPDATE personnel SET first_name=?,last_name=?,full_name=?,organizational_unit=?,position_code=?,
                       position_title=?,employment_group=?,employment_subtype=?,status=?,activity_area=?,actual_location=?,
                       company=?,chart_node_id=?,chart_page_no=?,extra_json=?,row_version=?,updated_at=?,updated_by=?
                       WHERE id=? AND row_version=?""",
                    (full_name, _text(old["last_name"]), full_name,
                     current_values["organizational_unit"], current_values["position_code"], current_values["position_title"],
                     current_values["employment_group"], current_values["employment_subtype"], current_values["status"],
                     current_values["activity_area"], current_values["actual_location"], current_values["company"],
                     chart_node, chart_page, canonical(extra), new_version, now, actor_id, person_id, old["row_version"]),
                )
                if conn.execute("SELECT changes()").fetchone()[0] != 1:
                    raise ValueError(f"Personnel {no}: رکورد هم‌زمان تغییر کرده است؛ Import لغو شد.")
                projection_values = {
                    "organizational_unit": current_values["organizational_unit"],
                    "position_code": current_values["position_code"],
                    "position_title": current_values["position_title"],
                    "actual_location": current_values["actual_location"],
                }
                from_id, to_id = repo._sync_person_projection(conn, person_id, projection_values, actor_id, now,
                                                               effective_date=effective_date,
                                                               assignment_id=("assignment-" + secrets.token_hex(16)) if assignment else None)
                after = dict(conn.execute("SELECT * FROM personnel WHERE id=?", (person_id,)).fetchone())
                repo._record(conn, actor_id, "update", "personnel", person_id, before, after, new_version)
                if item.get("changed_fields"):
                    updated_people += 1
                if assignment:
                    movement_id = "movement-" + secrets.token_hex(16)
                    movement_type = "position_change" if assignment["assignment_type"] == "approved_transfer" else "transfer"
                    conn.execute(
                        """INSERT INTO personnel_movements
                           (id,person_id,movement_type,effective_date,order_no,order_date,reason,note,
                            from_assignment_id,to_assignment_id,before_json,after_json,created_at,created_by,row_version)
                           VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,1)""",
                        (movement_id, person_id, movement_type, effective_date, "", "",
                         _text(assignment.get("reason")) or "به‌روزرسانی ماهانه Excel", _text(item.get("note")),
                         from_id, to_id, canonical(before), canonical(after), now, actor_id),
                    )
                    repo._record(conn, actor_id, "movement", "personnel_movement", movement_id, before, after, 1)
                    repo._record(conn, actor_id, "movement_update", "personnel", person_id, before, after, new_version)

            if assignment:
                conn.execute(
                    """INSERT INTO ui_monthly_assignments(person_id,personnel_no,assignment_type,target_page,target_node_id,
                       reason,batch_id,updated_at,updated_by) VALUES(?,?,?,?,?,?,?,?,?)
                       ON CONFLICT(person_id) DO UPDATE SET personnel_no=excluded.personnel_no,
                         assignment_type=excluded.assignment_type,target_page=excluded.target_page,
                         target_node_id=excluded.target_node_id,reason=excluded.reason,batch_id=excluded.batch_id,
                         updated_at=excluded.updated_at,updated_by=excluded.updated_by""",
                    (person_id, no, assignment["assignment_type"], int(assignment["target_page"]),
                     assignment["target_node_id"], _text(assignment.get("reason")), batch_id, now, actor_id),
                )
                assignments += 1

        summary = {
            "new_people": new_people,
            "updated_people": updated_people,
            "assignments": assignments,
            "backup": backup_filename,
            "sha256": str(plan.get("sha256", "")),
            "people_after": int(conn.execute("SELECT COUNT(*) FROM personnel").fetchone()[0]),
        }
        conn.execute(
            """INSERT INTO import_batches(id,source_name,source_kind,mode,row_count,accepted_count,
               warning_count,error_count,summary_json,created_by,created_at)
               VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
            (batch_id, str(plan.get("filename", "monthly.xlsx")), "monthly-xlsx", "production-apply",
             len(plan.get("people", [])), new_people + updated_people + assignments,
             int((plan.get("summary") or {}).get("warnings", 0)), 0, canonical(summary), actor_id, now),
        )
        repo._record(conn, actor_id, "monthly_import", "import_batch", batch_id, None, summary, 1)

    return {"batch_id": batch_id, **summary}
